#!/usr/bin/env python3
"""Gera dados/MODELO_apontamentos.xlsx - o modelo para pedir o export do ERP,
e dados/apontamentos_EXEMPLO.xlsx - dados ficticios so para ver a tela funcionando.

Rode: python gerar_modelo_apontamentos.py
"""

import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import mestre

RAIZ = Path(__file__).resolve().parent
DADOS = RAIZ / "dados"
DADOS.mkdir(exist_ok=True)

COLUNAS = ["Data", "Hora Inicial", "Hora Final", "HM", "Item", "Processo", "Recurso", "Operador", "Horas"]
AJUDA = [
    ("Data", "Dia do apontamento. dd/mm/aaaa. OBRIGATORIA."),
    ("Hora Inicial", "Hora em que comecou, hh:mm. Sem ela o programa nao consegue montar a grade hora a hora."),
    ("Hora Final", "Hora em que terminou, hh:mm. Se vier junto com a inicial, o programa calcula as horas sozinho."),
    ("HM", "Numero da HM. Pode vir como 'HM 024', 'HM024', '2026-024' ou so '024' - o programa acha o numero."),
    ("Item", "Qual ferramental dentro da HM. Pode ser so o codigo: CX02B, CX01A, MD01. Se vier vazio, "
             "a hora entra no setor sem quebra por item."),
    ("Processo", "Nome do processo no ERP: CNC GRANDE, CNC MEDIO, AJUSTAGEM, CAM, PROJETOS, SOLDAGEM... "
                 "A traducao para o setor do cronograma esta em base/de_para_erp_setor.csv. OBRIGATORIA."),
    ("Recurso", "Maquina ou posto. Opcional."),
    ("Operador", "Quem apontou. Opcional."),
    ("Horas", "Horas apontadas. Se vier vazia, o programa calcula por Hora Inicial e Hora Final."),
]


def cabecalho(ws, colunas):
    ws.append(colunas)
    fill = PatternFill("solid", fgColor="12304F")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    larguras = [12, 13, 13, 11, 34, 18, 22, 18, 9]
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


# ------------------------------------------------------------------ modelo
wb = Workbook()
ws = wb.active
ws.title = "Apontamentos"
cabecalho(ws, COLUNAS)
exemplos = [
    ("27/08/2026", "07:00", "11:30", "HM 024", "CX02B", "CAM", "CAM - equipe", "A. SOUZA", 4.5),
    ("27/08/2026", "13:00", "17:00", "HM 024", "CX02B", "CNC GRANDE", "GREBO 1580", "ALEX DA SILVA", 4.0),
    ("28/08/2026", "04:45", "11:30", "HM 051", "MD01", "AJUSTAGEM", "BANCADA", "EMERSON", 6.75),
]
for e in exemplos:
    ws.append(list(e))

wa = wb.create_sheet("Como preencher")
wa.append(["Coluna", "O que e"])
for c in wa[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="12304F")
wa.column_dimensions["A"].width = 16
wa.column_dimensions["B"].width = 110
for nome, txt in AJUDA:
    wa.append([nome, txt])
    wa.cell(row=wa.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
wa.append([])
wa.append(["", "A ordem das colunas nao importa. Os nomes podem variar um pouco - o programa reconhece "
               "sinonimos (OS, Ordem de Servico, Colaborador, Tempo Realizado...). Se o seu ERP usar um nome "
               "que ele nao reconhece, escreva o nome real em mapeamento.json."])
wa.cell(row=wa.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
wa.append(["", "Salve o arquivo na pasta dados com um nome que comece com 'apontamentos'. "
               "Pode ser .xlsx ou .csv."])
wa.cell(row=wa.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
wb.save(DADOS / "MODELO_apontamentos.xlsx")
print("gravado: dados/MODELO_apontamentos.xlsx")


# ------------------------------------------------------------------ exemplo
random.seed(11)
HOJE = date(2026, 9, 12)
PROC_POR_SETOR = {
    "Projeto/CAD": ["PROJETOS"], "CAM": ["CAM"],
    "Usinagem": ["CNC GRANDE", "CNC MEDIO", "FRESA CNC", "TORNO CNC"],
    "Solda": ["SOLDAGEM"], "Bancada": ["AJUSTAGEM"], "Convencional": ["FRESA CONV"],
    "Digitalização": ["DIGITALIZACAO"],
}
OPERADORES = {
    "Projeto/CAD": ["TIAGO POLAQUINI", "JOSE ESPINDOLA"], "CAM": ["JAISON", "ANDRE", "LUCIANO"],
    "Usinagem": ["ALEX DA SILVA", "MAICON", "GUSTAVO", "DANIEL SEVERINO", "JONATHAN"],
    "Solda": ["TIAGO", "HAROLD"], "Bancada": ["EMERSON", "AMARAL"],
    "Convencional": ["RAFAEL", "LEANDRO"], "Digitalização": ["NELSON JUNIOR"],
}

d = mestre.gerar()
linhas = []
for a in d["atividades"]:
    if a["hm"] not in ("HM 024", "HM 051") or not a["por_dia"]:
        continue
    procs = PROC_POR_SETOR.get(a["setor"], [a["setor"].upper()])
    ops = OPERADORES.get(a["setor"], ["OPERADOR"])
    # codigo curto do item, como o ERP costuma trazer
    cod = ""
    for p in str(a["item"]).replace("º", " ").split():
        if p.upper().startswith(("CX", "MD")) and any(ch.isdigit() for ch in p):
            cod = p.upper().strip(".")
    ritmo = random.choice([0.75, 0.9, 1.0, 1.0, 1.15])
    for dia in a["por_dia"]:
        dt = date.fromisoformat(dia["data"])
        if dt > HOJE:
            continue
        restante = dia["horas"] * ritmo * random.uniform(0.85, 1.15)
        hora = 4 if "TURNO1" in (a["calendarios"] or []) else 7
        while restante > 0.2 and hora < 22:
            if hora in (11, 12):
                hora = 13
                continue
            bloco = min(restante, random.choice([1.5, 2.0, 2.0, 3.0]))
            ini = f"{hora:02d}:00"
            fim_m = int(hora * 60 + bloco * 60)
            linhas.append([dt.strftime("%d/%m/%Y"), ini, f"{fim_m // 60:02d}:{fim_m % 60:02d}",
                           a["hm"], cod, random.choice(procs),
                           a["recurso"][:22], random.choice(ops), round(bloco, 2)])
            restante -= bloco
            hora += int(bloco) + 1

linhas.sort(key=lambda r: (datetime.strptime(r[0], "%d/%m/%Y"), r[1]))
wb2 = Workbook()
ws2 = wb2.active
ws2.title = "Apontamentos"
cabecalho(ws2, COLUNAS)
for l in linhas:
    ws2.append(l)
wb2.save(DADOS / "apontamentos_EXEMPLO.xlsx")
print(f"gravado: dados/apontamentos_EXEMPLO.xlsx ({len(linhas)} lancamentos, "
      f"{sum(l[8] for l in linhas):.1f} h)")
