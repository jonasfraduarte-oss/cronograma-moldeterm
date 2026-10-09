#!/usr/bin/env python3
"""
IMPORTADOR DAS PLANILHAS DE HM (modelo Tupy "HM JOI ...")
========================================================
Le todos os .xlsx da pasta pedidos/hms e extrai, de cada um:
  - os ITENS da HM (aba "RFQ Summary"): servico + descricao + requisicao/linha
  - as HORAS por processo de cada item (aba "Orcamento Hora")

Grava o resultado em pedidos/horas_itens.csv, que e o que alimenta o nivel de
ITEM do cronograma.

Regras:
  - a coluna de horas usada e "Quantidade Horas Pos Kickoff"; quando ela esta
    vazia, usa "Quantidade Horas Relatorio Final" (a coluna preenchida manda);
  - Material P20, Material 1045 e Acessorios NAO sao horas e ficam de fora;
  - nada e estimado: bloco sem horas vira pendencia, nao virou numero.

Uso:  python importar_hms.py
"""

import csv
import re
import unicodedata
from pathlib import Path

RAIZ = Path(__file__).resolve().parent
PASTA_HMS = RAIZ / "pedidos" / "hms"
SAIDA = RAIZ / "pedidos" / "horas_itens.csv"
DE_PARA = RAIZ / "base" / "de_para_processos.csv"

SERVICOS = ("DESIGN", "ALTERATION", "SCANNING", "REPAIR", "REPARO")


def nm(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().upper()


def carregar_de_para():
    mapa = {}
    if DE_PARA.exists():
        texto = DE_PARA.read_text(encoding="utf-8-sig")
        for r in csv.DictReader(texto.splitlines(), delimiter=";"):
            mapa[nm(r["processo_planilha_hm"])] = {
                "setor": r["setor_cronograma"].strip(),
                "origem": r.get("origem", ""), "obs": r.get("observacao", ""),
            }
    return mapa


def numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(".", "").replace(",", "."))
    except ValueError:
        return None


def data_br(v):
    """Aceita data do Excel ou texto. Devolve (iso, aviso)."""
    if v is None or str(v).strip() == "":
        return None, ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d"), ""
    t = str(v).strip()
    from datetime import datetime as dt
    for f in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.strptime(t, f).strftime("%Y-%m-%d"), ""
        except ValueError:
            pass
    return None, f"data '{t}' nao esta num formato reconhecido"


def ler_cabecalho(wb):
    """Tira da aba RFQ Summary: cliente, PN, projeto e as datas TARGET."""
    if "RFQ Summary" not in wb.sheetnames:
        return {}, []
    ws = wb["RFQ Summary"]
    linhas = [list(r) for r in ws.iter_rows(min_row=1, max_row=10, values_only=True)]
    cab, avisos = {}, []
    for l in linhas:
        v = [("" if c is None else c) for c in l]
        txts = [str(x).strip() for x in v]
        if any(nm(x) == "TUPY JOINVILLE" for x in txts):
            nao_vazios = [x for x in txts if x]
            if len(nao_vazios) > 2:
                cab["cliente"] = nao_vazios[2] if len(nao_vazios) > 2 else ""
                cab["pn"] = nao_vazios[3] if len(nao_vazios) > 3 else ""
                cab["projeto"] = nao_vazios[4] if len(nao_vazios) > 4 else ""
        rot = next((nm(x) for x in txts if nm(x) in ("FINAL", "INICAL", "INICIAL")), None)
        if rot:
            datas = []
            achou = False
            for c in v:
                if not achou:
                    if isinstance(c, str) and nm(c) == rot:
                        achou = True
                    continue
                if c is None or str(c).strip() == "":
                    continue
                iso, aviso = data_br(c)
                if aviso:
                    avisos.append(aviso)
                if iso:
                    datas.append(iso)
            chave = "final" if rot == "FINAL" else "inicial"
            if datas:
                cab[f"{chave}_tupy"] = datas[0]
                if len(datas) > 1:
                    cab[f"{chave}_fornecedor"] = datas[1]
    return cab, avisos


def ler_hm(caminho, mapa):
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True)
    m = re.search(r"2026[-_](\d{3})", caminho.name)
    hm = f"HM {m.group(1)}" if m else caminho.stem
    pn = None
    m2 = re.search(r"-_?(\d{2}-\d{2}-\d{3})", caminho.name)
    if m2:
        pn = m2.group(1)
    cab_rfq, avisos_data = ler_cabecalho(wb)

    # ---------- itens declarados
    itens_rfq = []
    if "RFQ Summary" in wb.sheetnames:
        for l in wb["RFQ Summary"].iter_rows(values_only=True):
            v = [("" if c is None else str(c).strip()) for c in l]
            if len(v) > 8 and v[2] and v[5] and nm(v[5]) in SERVICOS:
                itens_rfq.append({"num": v[2], "requisicao": v[3], "linha": v[4],
                                  "servico": v[5], "descricao": v[7]})

    # ---------- horas por bloco
    aba = next((s for s in wb.sheetnames if "amento" in nm(s) or "ORCAMENTO" in nm(s)), None)
    linhas, sem_horas = [], []
    if aba:
        bloco = None
        for l in wb[aba].iter_rows(values_only=True):
            v = list(l) + [None] * 10
            if isinstance(v[1], str) and nm(v[1]) in SERVICOS and v[3]:
                if bloco and not bloco["achou"]:
                    sem_horas.append(bloco["rotulo"])
                bloco = {"servico": v[1].strip(), "descricao": str(v[3]).strip(),
                         "rotulo": f"{v[1].strip()} - {str(v[3]).strip()}", "achou": False}
                continue
            if bloco and isinstance(v[2], str):
                proc = v[2].strip()
                chave = nm(proc)
                if chave not in mapa:
                    continue
                setor = mapa[chave]["setor"]
                h = numero(v[3])
                if not h:
                    h = numero(v[4])
                if not h or h <= 0:
                    continue
                if setor in ("-", ""):
                    continue  # material, nao e hora
                bloco["achou"] = True
                linhas.append({
                    "hm": hm, "pn": pn or "", "item": bloco["descricao"], "servico": bloco["servico"],
                    "processo": proc, "setor": setor, "horas_base": round(h, 2),
                    "origem": "PLANILHA HM", "observacao": mapa[chave]["obs"],
                })
        if bloco and not bloco["achou"]:
            sem_horas.append(bloco["rotulo"])

    cab_rfq["hm"] = hm
    cab_rfq["pn"] = cab_rfq.get("pn") or pn or ""
    cab_rfq["arquivo"] = caminho.name
    cab_rfq["avisos"] = " | ".join(avisos_data)
    return hm, pn, itens_rfq, linhas, sem_horas, cab_rfq


def main():
    mapa = carregar_de_para()
    if not PASTA_HMS.exists():
        print(f"Pasta {PASTA_HMS} nao encontrada.")
        return
    todas, avisos, cabecalhos = [], [], []
    for arq in sorted(PASTA_HMS.glob("*.xlsx")):
        if arq.name.startswith("~$"):
            continue
        hm, pn, itens, linhas, sem, cab = ler_hm(arq, mapa)
        todas.extend(linhas)
        cabecalhos.append(cab)
        if cab.get("avisos"):
            avisos.append(f"{hm}: {cab['avisos']}")
        tot = sum(l["horas_base"] for l in linhas)
        print(f"{hm} ({pn or '-'})  {len(itens)} itens declarados | {len(linhas)} lancamentos | {tot:.2f} h base")
        for s in sem:
            avisos.append(f"{hm}: bloco '{s}' esta sem horas preenchidas na planilha")
        if itens and not linhas:
            avisos.append(f"{hm}: a planilha tem itens declarados mas NENHUMA hora preenchida")

    todas.sort(key=lambda r: (r["hm"], r["item"], r["servico"], r["setor"]))
    with open(SAIDA, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["hm", "pn", "item", "servico", "processo", "setor",
                                          "horas_base", "origem", "observacao"], delimiter=";")
        w.writeheader()
        w.writerows(todas)
    campos_cab = ["hm", "arquivo", "cliente", "pn", "projeto", "inicial_tupy", "inicial_fornecedor",
                  "final_tupy", "final_fornecedor", "avisos"]
    with open(RAIZ / "pedidos" / "hms_cabecalho.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=campos_cab, delimiter=";", extrasaction="ignore")
        w.writeheader()
        for c in sorted(cabecalhos, key=lambda x: x.get("hm", "")):
            w.writerow({k: c.get(k, "") for k in campos_cab})

    print(f"\n{len(todas)} linhas gravadas em {SAIDA.relative_to(RAIZ)}")
    print(f"{len(cabecalhos)} cabecalhos gravados em pedidos/hms_cabecalho.csv")
    if avisos:
        print("\nAvisos:")
        for a in avisos:
            print("  !", a)


if __name__ == "__main__":
    main()
